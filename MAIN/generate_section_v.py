import os
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_section_v_doc():
    doc = docx.Document()

    # Apply Normal style formatting
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Times New Roman'
    font_normal.size = Pt(10.5)

    def add_justified_paragraph(text):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    # Heading
    heading = doc.add_paragraph()
    heading_run = heading.add_run("V. PERFORMANCE EVALUATION & DISCUSSION")
    heading_run.bold = True
    heading_run.font.name = 'Times New Roman'
    heading_run.font.size = Pt(12)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Introduction body text
    add_justified_paragraph("To evaluate the viability of the relay architecture, experimental tests were conducted using standard file sizes (10MB, 100MB, 1GB) to simulate typical usage scenarios, including document syncing and media streaming.")
    
    doc.add_paragraph()

    # Table for Average Upload Latency
    doc.add_paragraph("Average Upload Latency (Local LAN vs. Internet Relay):").runs[0].bold = True
    
    table_latency = doc.add_table(rows=1, cols=3)
    table_latency.style = 'Table Grid'
    hdr_cells = table_latency.rows[0].cells
    hdr_cells[0].text = 'File Size'
    hdr_cells[1].text = 'Local LAN Latency'
    hdr_cells[2].text = 'Internet Relay Latency'
    
    for cell in hdr_cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                
    latency_data = [
        ('10MB', '~150ms', '~2.5s'),
        ('100MB', '~1.5s', '~22s'),
        ('1GB', '~12s', '~4.5 mins')
    ]
    for size, lan, relay in latency_data:
        row_cells = table_latency.add_row().cells
        row_cells[0].text = size
        row_cells[1].text = lan
        row_cells[2].text = relay

    doc.add_paragraph()

    # Bullet points for Hardware Utilization
    doc.add_paragraph("Desktop Hardware Utilization (Tested on standard i5/i7 H-series processor):").runs[0].bold = True
    doc.add_paragraph("Idle State: CPU < 1%, Memory ~450MB", style='List Bullet')
    doc.add_paragraph("Active MinIO Read/Write (Chunking/Hashing): CPU 18-24%, Memory Peak ~1.1GB", style='List Bullet')
    
    doc.add_paragraph()
    
    # Bullet points for Mobile Battery Impact
    doc.add_paragraph("Mobile Battery Impact (Flutter Client):").runs[0].bold = True
    doc.add_paragraph("Active background synchronization of 1GB payload: ~3% battery drain.", style='List Bullet')
    doc.add_paragraph("Idle background polling: < 0.5% per hour.", style='List Bullet')
    
    doc.add_paragraph()
    
    # Conclusion body text
    add_justified_paragraph("An analytical comparison of these metrics against direct peer-to-peer and centralized setups illustrates the acceptable overhead introduced by the relay. While centralized platforms benefit from edge caching, our relay model maintains competitive latency without sacrificing user privacy.")

    output_path = os.path.abspath("Section_V_Performance.docx")
    doc.save(output_path)
    print(f"File successfully created at: {output_path}")

if __name__ == "__main__":
    create_section_v_doc()
